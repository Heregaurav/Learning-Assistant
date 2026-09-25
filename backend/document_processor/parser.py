import re
import tempfile
from functools import lru_cache
from pathlib import Path
from typing import Optional

from pydantic import BaseModel, Field


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".pptx"}
ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "text/plain",
    "text/markdown",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/octet-stream",
}
MAX_FILE_SIZE_BYTES = 12 * 1024 * 1024
MAX_OCR_PAGES = 20


class ParsedDocument(BaseModel):
    title: str
    text: str
    extension: str
    mime_type: str
    page_count: Optional[int] = None
    word_count: int = 0
    source_name: str
    warnings: list[str] = Field(default_factory=list)


class UnsupportedDocumentError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


def safe_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def detect_mime_type(filename: str, sniff: Optional[str] = None) -> str:
    suffix = safe_extension(filename)
    if suffix == ".pdf":
        return "application/pdf"
    if suffix in {".docx", ".pptx"}:
        return (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if suffix == ".docx"
            else "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
    if suffix in {".txt", ".md", ".markdown"}:
        return "text/plain" if suffix == ".txt" else "text/markdown"
    return sniff or "application/octet-stream"


def validate_upload(filename: str, content_type: str, size_bytes: int) -> None:
    ext = safe_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedDocumentError(
            "Unsupported file type. Please upload a PDF, DOCX, TXT, Markdown, or PPTX document."
        )
    if not content_type:
        content_type = detect_mime_type(filename)
    mime = content_type.split(";", 1)[0].strip().lower()
    if mime not in ALLOWED_MIME_TYPES and ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedDocumentError("This file type is not supported.")
    if size_bytes <= 0:
        raise EmptyDocumentError("The uploaded file is empty.")
    if size_bytes > MAX_FILE_SIZE_BYTES:
        raise ValueError("This document is too large. Please upload a file under 12 MB.")


def _clean_text(raw: str) -> str:
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\t", " ")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{2,}", "\n\n", text)
    text = text.strip()
    return text


def _extract_from_txt(path: Path) -> str:
    try:
        return _clean_text(path.read_text(encoding="utf-8"))
    except UnicodeDecodeError:
        return _clean_text(path.read_text(encoding="latin-1"))


def _extract_from_docx(path: Path) -> str:
    try:
        from docx import Document

        document = Document(str(path))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        return _clean_text("\n\n".join(parts))
    except Exception:
        return ""


def _extract_from_pptx(path: Path) -> str:
    try:
        import zipfile
        import re as regex

        with zipfile.ZipFile(path) as zf:
            text_parts = []
            for name in zf.namelist():
                if name.endswith(".xml") and "ppt/slides/slide" in name:
                    xml = zf.read(name).decode("utf-8", errors="ignore")
                    texts = regex.findall(r">([^<>]+)</a:t>", xml)
                    if texts:
                        text_parts.extend(texts)
            return _clean_text("\n\n".join(text_parts))
    except Exception:
        return ""


def _extract_from_pdf(path: Path) -> str:
    try:
        import pypdf

        reader = pypdf.PdfReader(str(path))
        page_texts = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text and page_text.strip():
                page_texts.append(page_text)
        text = "\n\n".join(page_texts)
        return _clean_text(text)
    except Exception:
        return ""


@lru_cache(maxsize=1)
def _get_ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


def _ocr_pdf(path: Path) -> tuple[str, int]:
    pdf = None
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(path))
        page_count = len(pdf)
        page_texts = []
        engine = _get_ocr_engine()
        for page_index in range(min(page_count, MAX_OCR_PAGES)):
            page = pdf[page_index]
            bitmap = page.render(scale=2.0)
            try:
                result, _ = engine(bitmap.to_numpy())
            finally:
                bitmap.close()
                page.close()
            lines = [
                str(line[1]).strip()
                for line in (result or [])
                if len(line) > 2 and float(line[2]) >= 0.3 and str(line[1]).strip()
            ]
            if lines:
                page_texts.append("\n".join(lines))
        return _clean_text("\n\n".join(page_texts)), page_count
    except Exception as exc:
        raise ValueError("OCR could not read this PDF. Check the backend OCR dependencies and try again.") from exc
    finally:
        if pdf is not None:
            pdf.close()


def parse_document(filename: str, content_type: str, file_bytes: bytes, *, source_name: str | None = None) -> ParsedDocument:
    validate_upload(filename, content_type, len(file_bytes))
    ext = safe_extension(filename)
    if not file_bytes.strip():
        raise EmptyDocumentError("The uploaded document is empty.")

    suffix = ext.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(file_bytes)
        tmp_path = Path(tmp.name)

    try:
        page_count = None
        warnings: list[str] = []
        if suffix == ".txt":
            text = _extract_from_txt(tmp_path)
        elif suffix == ".md" or suffix == ".markdown":
            text = _extract_from_txt(tmp_path)
        elif suffix == ".docx":
            text = _extract_from_docx(tmp_path)
        elif suffix == ".pptx":
            text = _extract_from_pptx(tmp_path)
        elif suffix == ".pdf":
            text = _extract_from_pdf(tmp_path)
            if not text.strip():
                text, page_count = _ocr_pdf(tmp_path)
                if page_count > MAX_OCR_PAGES:
                    warnings.append(
                        f"OCR processed the first {MAX_OCR_PAGES} of {page_count} pages."
                    )
        else:
            text = _extract_from_txt(tmp_path)

        if not text or not text.strip():
            raise ValueError(
                "No readable text was detected in this document, even after OCR."
            )

        title = Path(filename).stem.strip() or "Document"
        title = re.sub(r"[_-]+", " ", title)
        words = re.findall(r"\b\w+\b", text)
        return ParsedDocument(
            title=title,
            text=text,
            extension=suffix,
            mime_type=detect_mime_type(filename, content_type),
            page_count=page_count,
            word_count=len(words),
            source_name=source_name or filename,
            warnings=warnings,
        )
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass
